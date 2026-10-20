"""Small helpers on top of python-docx for the HRX design & technical document:
Raleway everywhere, the app's colours, tables in the app's style, a table of contents
field, and the Raleway font embedded in the .docx so it shows on machines without it.
"""
import io
import os
import re
import uuid
import zipfile

from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Cm

FONT = "Raleway"

# The app's colour tokens (hrx2026/webapp/css/style.css :root)
INK = "33343A"
TXT2 = "6D6E72"
TXT3 = "9B9CA1"
ICE = "26AAE2"
ICE_DEEP = "1888BC"
VIOLET = "756EE5"
MAGENTA = "9B39B0"
FROST = "F3F5F7"
LINE = "E3E6EA"
MINT = "1FAE72"
AMBER = "D97706"
AMBER_SOFT = "FEF3C7"
SIGNAL = "E5484D"


def rgb(hex_colour):
    return RGBColor.from_string(hex_colour)


def set_run_font(run, size=None, bold=None, colour=None, italic=None):
    run.font.name = FONT
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), FONT)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic
    if colour is not None:
        run.font.color.rgb = rgb(colour)
    return run


def style_document(document):
    """Raleway and the app's ink colour for every style the document uses."""
    styles = document.styles
    for name in ["Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3", "List Bullet", "List Number", "Caption",
                 "TOC Heading", "toc 1", "toc 2"]:
        try:
            style = styles[name]
        except KeyError:
            continue
        rfonts = style.element.get_or_add_rPr().get_or_add_rFonts()
        for attribute in ["w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"]:
            rfonts.set(qn(attribute), FONT)
        for attribute in ["w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"]:
            if rfonts.get(qn(attribute)) is not None:
                del rfonts.attrib[qn(attribute)]
        style.font.name = FONT

    defaults = document.styles.element.find(qn("w:docDefaults"))
    if defaults is not None:
        rpr = defaults.find(qn("w:rPrDefault") + "/" + qn("w:rPr"))
        if rpr is not None:
            for old in rpr.findall(qn("w:rFonts")):
                rpr.remove(old)
            rfonts = OxmlElement("w:rFonts")
            for attribute in ["w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"]:
                rfonts.set(qn(attribute), FONT)
            rpr.insert(0, rfonts)

    # Front matter headings (document control, contents): like Heading 1, but kept out
    # of the table of contents.
    from docx.enum.style import WD_STYLE_TYPE
    front = styles.add_style("HRX Front Heading", WD_STYLE_TYPE.PARAGRAPH)
    front.base_style = styles["Normal"]
    front.font.size = Pt(18)
    front.font.bold = True
    front.font.color.rgb = rgb(INK)
    front.paragraph_format.space_before = Pt(18)
    front.paragraph_format.space_after = Pt(8)
    front.paragraph_format.page_break_before = True
    front.paragraph_format.keep_with_next = True
    paragraph_border(front.element.get_or_add_pPr(), "bottom", ICE, 8)

    # The contents' own entry styles (the default template only lists them as latent)
    for level, indent, size, bold in [(1, 0, 10, True), (2, 0.6, 9.5, False)]:
        name = "toc " + str(level)
        if name not in [style.name for style in styles]:
            toc = styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
            toc.base_style = styles["Normal"]
            toc.font.size = Pt(size)
            toc.font.bold = bold
            toc.paragraph_format.left_indent = Cm(indent)
            toc.paragraph_format.space_before = Pt(6 if level == 1 else 0)
            toc.paragraph_format.space_after = Pt(2)

    normal = styles["Normal"]
    normal.font.size = Pt(10)
    normal.font.color.rgb = rgb(INK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.2

    for name, size, colour, before, after in [("Heading 1", 18, INK, 18, 8), ("Heading 2", 13, ICE_DEEP, 14, 4), ("Heading 3", 11, INK, 10, 3)]:
        style = styles[name]
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.italic = False
        style.font.color.rgb = rgb(colour)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
    # Each top-level section starts on a new page, under a thin ice-blue rule
    heading1 = styles["Heading 1"]
    heading1.paragraph_format.page_break_before = True
    paragraph_border(heading1.element.get_or_add_pPr(), "bottom", ICE, 8)

    for name in ["List Bullet", "List Number"]:
        styles[name].paragraph_format.space_after = Pt(3)


def paragraph_border(ppr, side, colour, size_eighths):
    borders = ppr.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        ppr.append(borders)
    edge = OxmlElement("w:" + side)
    edge.set(qn("w:val"), "single")
    edge.set(qn("w:sz"), str(size_eighths))
    edge.set(qn("w:space"), "4")
    edge.set(qn("w:color"), colour)
    borders.append(edge)


def shade(cell, fill):
    tcpr = cell._element.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), fill)
    tcpr.append(shading)


def cell_margins(table, top=60, bottom=60, left=100, right=100):
    tblpr = table._element.tblPr
    margins = OxmlElement("w:tblCellMar")
    for side, value in [("top", top), ("left", left), ("bottom", bottom), ("right", right)]:
        edge = OxmlElement("w:" + side)
        edge.set(qn("w:w"), str(value))
        edge.set(qn("w:type"), "dxa")
        margins.append(edge)
    tblpr.append(margins)


def table_borders(table, colour=LINE):
    tblpr = table._element.tblPr
    borders = OxmlElement("w:tblBorders")
    for side in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        edge = OxmlElement("w:" + side)
        edge.set(qn("w:val"), "single")
        edge.set(qn("w:sz"), "4")
        edge.set(qn("w:space"), "0")
        edge.set(qn("w:color"), colour)
        borders.append(edge)
    tblpr.append(borders)


def fix_widths(table, widths_cm):
    """Gives a table a fixed layout and a column grid, so Word and LibreOffice both
    keep the column widths (cell widths alone are not enough for LibreOffice)."""
    from docx.shared import Cm as _Cm
    tbl = table._tbl
    tblpr = tbl.tblPr
    total = sum(widths_cm)
    width = OxmlElement("w:tblW")
    width.set(qn("w:w"), str(int(_Cm(total).twips)))
    width.set(qn("w:type"), "dxa")
    for old in tblpr.findall(qn("w:tblW")):
        tblpr.remove(old)
    tblpr.append(width)
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblpr.append(layout)
    grid = tbl.tblGrid
    for column in list(grid):
        grid.remove(column)
    for value in widths_cm:
        column = OxmlElement("w:gridCol")
        column.set(qn("w:w"), str(int(_Cm(value).twips)))
        grid.append(column)
    for row in table.rows:
        for index, value in enumerate(widths_cm):
            if index < len(row.cells):
                row.cells[index].width = _Cm(value)


def repeat_header(row):
    trpr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    trpr.append(header)


def write_cell(cell, text, size=9, bold=False, colour=INK, code=False):
    """Writes text into a cell. **bold** and `code` spans are honoured."""
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.1
    for line_index, line in enumerate(str(text).split("\n")):
        if line_index:
            paragraph = cell.add_paragraph()
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.1
        add_rich(paragraph, line, size=size, bold=bold, colour=colour)


def add_rich(paragraph, text, size=10, bold=False, colour=INK):
    """Adds text with **bold** and `code` spans to a paragraph."""
    for part in re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text):
        if not part:
            continue
        if part.startswith("**"):
            set_run_font(paragraph.add_run(part[2:-2]), size=size, bold=True, colour=colour)
        elif part.startswith("`"):
            run = paragraph.add_run(part[1:-1])
            set_run_font(run, size=size - 0.5, bold=bold, colour=ICE_DEEP)
        else:
            set_run_font(paragraph.add_run(part), size=size, bold=bold, colour=colour)
    return paragraph


def add_table(document, header, rows, widths_cm, size=8.5, header_fill=FROST):
    table = document.add_table(rows=1, cols=len(header))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table_borders(table)
    cell_margins(table)
    for index, title in enumerate(header):
        cell = table.rows[0].cells[index]
        write_cell(cell, title.upper(), size=7.5, bold=True, colour=TXT2)
        shade(cell, header_fill)
    repeat_header(table.rows[0])
    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            write_cell(cells[index], value, size=size)
    fix_widths(table, widths_cm)
    document.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_toc(document, entries=None):
    """A table of contents field (headings 1-2). entries: [(level, text, page)] shown as
    the field's last result, so the document opens with a filled-in contents; Word
    refreshes it with Update Field."""
    paragraph = document.add_paragraph()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    begin.set(qn("w:dirty"), "false")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = ' TOC \\o "1-2" \\h \\z \\u '
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    for element in [begin, instruction, separate]:
        run = paragraph.add_run()
        run._r.append(element)
    last = paragraph
    if entries:
        for index, (level, text, page) in enumerate(entries):
            entry = paragraph if index == 0 else document.add_paragraph()
            entry.style = document.styles["toc " + str(level)]
            entry.paragraph_format.tab_stops.add_tab_stop(Cm(16.4), alignment=2, leader=1)  # right, dots
            set_run_font(entry.add_run(text + "\t" + str(page)), size=10 if level == 1 else 9.5, bold=(level == 1))
            last = entry
    else:
        set_run_font(paragraph.add_run("Right-click and choose Update Field to show the contents."), size=9, colour=TXT2)
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run = last.add_run()
    run._r.append(end)


def add_field(paragraph, instruction, size=8, colour=TXT2):
    for kind, text in [("begin", None), ("instr", instruction), ("separate", None), ("text", "1"), ("end", None)]:
        run = paragraph.add_run()
        set_run_font(run, size=size, colour=colour)
        if kind == "instr":
            element = OxmlElement("w:instrText")
            element.set(qn("xml:space"), "preserve")
            element.text = " " + text + " "
            run._r.append(element)
        elif kind == "text":
            run.text = text
        else:
            element = OxmlElement("w:fldChar")
            element.set(qn("w:fldCharType"), kind)
            run._r.append(element)


def page_break(document):
    document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


# ── Font embedding (ECMA-376 Part 1, 17.8.1: obfuscated .odttf parts) ──

# Which characters of the "{XXXXXXXX-XXXX-XXXX-XXXX-XXXXXXXXXXXX}" key make each key
# byte, as Word and LibreOffice read it.
KEY_POSITIONS = [35, 33, 31, 29, 27, 25, 22, 20, 17, 15, 12, 10, 7, 5, 3, 1]


def obfuscate(font_bytes, key):
    key_bytes = [int(key[position:position + 2], 16) for position in KEY_POSITIONS]
    data = bytearray(font_bytes)
    for index in range(16):
        data[index] ^= key_bytes[index]
        data[index + 16] ^= key_bytes[index]
    return bytes(data)


def embed_fonts(docx_path, font_files):
    """font_files: {"embedRegular": path, "embedBold": path, ...} for Raleway."""
    with zipfile.ZipFile(docx_path) as source:
        parts = {name: source.read(name) for name in source.namelist()}

    rels_name = "word/_rels/fontTable.xml.rels"
    rels = parts.get(rels_name, b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
                     b'<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"></Relationships>').decode("utf-8")
    font_xml = ['<w:font w:name="Raleway"><w:panose1 w:val="020B0503030101060003"/><w:charset w:val="00"/>'
                '<w:family w:val="swiss"/><w:pitch w:val="variable"/>']
    for index, (element, path) in enumerate(font_files.items(), start=1):
        key = "{" + str(uuid.uuid4()).upper() + "}"
        part = "fonts/raleway" + str(index) + ".odttf"
        with open(path, "rb") as file:
            parts["word/" + part] = obfuscate(file.read(), key)
        rel_id = "rIdRaleway" + str(index)
        rels = rels.replace("</Relationships>", '<Relationship Id="' + rel_id + '" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/font" Target="' + part + '"/></Relationships>')
        font_xml.append('<w:' + element + ' r:id="' + rel_id + '" w:fontKey="' + key + '"/>')
    font_xml.append("</w:font>")
    parts[rels_name] = rels.encode("utf-8")

    font_table = parts["word/fontTable.xml"].decode("utf-8")
    font_table = re.sub(r'<w:font w:name="Raleway">.*?</w:font>', "", font_table, flags=re.S)
    if 'xmlns:r="' not in font_table.split(">", 2)[1]:
        font_table = font_table.replace("<w:fonts ", '<w:fonts xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" ', 1)
    font_table = font_table.replace("</w:fonts>", "".join(font_xml) + "</w:fonts>")
    parts["word/fontTable.xml"] = font_table.encode("utf-8")

    types = parts["[Content_Types].xml"].decode("utf-8")
    if 'Extension="odttf"' not in types:
        types = types.replace("<Default ", '<Default Extension="odttf" ContentType="application/vnd.openxmlformats-officedocument.obfuscatedFont"/><Default ', 1)
    parts["[Content_Types].xml"] = types.encode("utf-8")

    # The theme's own fonts too, so nothing falls back to Calibri or Cambria
    if "word/theme/theme1.xml" in parts:
        theme = parts["word/theme/theme1.xml"].decode("utf-8")
        theme = re.sub(r'<a:latin typeface="[^"]*"', '<a:latin typeface="Raleway"', theme)
        parts["word/theme/theme1.xml"] = theme.encode("utf-8")

    settings = parts["word/settings.xml"].decode("utf-8")
    if "embedTrueTypeFonts" not in settings:
        settings = re.sub(r"(<w:settings[^>]*>)", r"\1<w:embedTrueTypeFonts/>", settings, count=1)
    parts["word/settings.xml"] = settings.encode("utf-8")

    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as target:
        target.writestr("[Content_Types].xml", parts.pop("[Content_Types].xml"))
        for name, data in parts.items():
            target.writestr(name, data)
    with open(docx_path, "wb") as file:
        file.write(buffer.getvalue())


def find_raleway():
    """The Raleway TTFs on this machine: {"embedRegular": path, ...}."""
    folders = [os.path.expanduser("~/.fonts"), os.path.expanduser("~/.local/share/fonts"), os.path.expanduser("~/Library/Fonts"),
               "/Library/Fonts", "/usr/share/fonts/truetype/raleway"]
    wanted = {"embedRegular": "Raleway-Regular.ttf", "embedBold": "Raleway-Bold.ttf", "embedItalic": "Raleway-Italic.ttf"}
    found = {}
    for element, name in wanted.items():
        for folder in folders:
            if os.path.exists(os.path.join(folder, name)):
                found[element] = os.path.join(folder, name)
                break
    return found
